#!/usr/bin/env python3
from __future__ import annotations

import base64
import json
import os
import re
import tempfile
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

SOURCE_URL = os.environ.get(
    "SCIENTIFIC_ONEDRIVE_URL",
    "https://1drv.ms/x/c/f0437e1066d439c2/IQAG-oAXUhAgR6GLr-icPyUEAV8yn94El0QRd_kLoOFteuE?e=D0fVWv",
)
OUTPUT = Path(os.environ.get("SCIENTIFIC_OUTPUT_JSON", "data/cientificas.json"))
MONTHS = {
    "Enero": 1, "Febrero": 2, "Marzo": 3, "Abril": 4, "Mayo": 5, "Junio": 6,
    "Julio": 7, "Agosto": 8, "Septiembre": 9, "Octubre": 10,
    "Noviembre": 11, "Diciembre": 12, "Dicimbre": 12,
}


def norm(v):
    return "" if v is None else str(v).strip()


def fetch_bytes(url: str, headers=None, data=None, timeout=120):
    req = urllib.request.Request(url, data=data, headers=headers or {})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read(), resp.geturl(), resp.headers.get("Content-Type", "")


def download(dest: Path):
    token_body = json.dumps({"appId": "5cbed6ac-a083-4e14-b191-b4ba07653de2"}).encode("utf-8")
    token_bytes, _, _ = fetch_bytes(
        "https://api-badgerp.svc.ms/v1.0/token",
        headers={
            "User-Agent": "Mozilla/5.0 EfemeridesCientificas/1.0",
            "Content-Type": "application/json",
            "Accept": "application/json",
        },
        data=token_body,
    )
    token_payload = json.loads(token_bytes.decode("utf-8"))
    badger_token = token_payload.get("token")
    if not badger_token:
        raise RuntimeError("OneDrive no devolvió el token temporal.")

    encoded = base64.urlsafe_b64encode(SOURCE_URL.encode("utf-8")).decode("ascii").rstrip("=")
    api_url = (
        "https://my.microsoftpersonalcontent.com/_api/v2.0/shares/"
        f"u!{encoded}/driveitem?select=@content.downloadUrl,name"
    )
    item_bytes, _, _ = fetch_bytes(
        api_url,
        headers={
            "User-Agent": "Mozilla/5.0 EfemeridesCientificas/1.0",
            "Accept": "application/json",
            "Prefer": "autoredeem",
            "Authorization": f"Badger {badger_token}",
        },
    )
    item = json.loads(item_bytes.decode("utf-8"))
    download_url = item.get("@content.downloadUrl")
    if not download_url:
        raise RuntimeError("OneDrive no devolvió una URL de descarga para el Excel científico.")

    data, _, content_type = fetch_bytes(
        download_url,
        headers={"User-Agent": "Mozilla/5.0 EfemeridesCientificas/1.0"},
    )
    if not data.startswith(b"PK\x03\x04"):
        raise RuntimeError(
            f"OneDrive no devolvió un XLSX científico válido ({content_type}, {len(data)} bytes)."
        )

    dest.write_bytes(data)
    if not dest.exists() or dest.stat().st_size < 1000:
        raise RuntimeError("Descarga científica vacía o incompleta.")


def build(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError("Excel científico vacío.")

    events, variable, undated, issues = [], [], [], []
    for idx, row in enumerate(rows[1:], start=2):
        if len(row) < 6:
            continue
        day, month_raw, year, name, desc, cat = row[:6]
        name = norm(name)
        if not name:
            continue

        month = MONTHS.get(norm(month_raw))
        yr = int(year) if isinstance(year, (int, float)) else None
        common = {
            "id": f"sci-{idx-1}",
            "name": name,
            "description": norm(desc),
            "category": norm(cat),
            "year": yr,
            "source_row": idx,
        }

        if isinstance(day, (int, float)) and month:
            events.append({**common, "day": int(day), "month": month})
        elif isinstance(day, str) and month and not re.match(r"^\d{4}-\d{2}-\d{2}", day):
            variable.append({**common, "month": month, "date_text": day.strip()})
        else:
            if hasattr(day, "date") and month:
                issues.append({
                    "source_row": idx,
                    "name": name,
                    "issue": f"Campo Día contiene {day.date().isoformat()} mientras Mes={norm(month_raw)}",
                })
            undated.append({
                **common,
                "date_text": norm(day) or None,
                "month_text": norm(month_raw) or None,
            })

    events.sort(key=lambda x: (x["month"], x["day"], x["name"].casefold()))
    return {
        "meta": {
            "title": "Efemérides científicas, académicas y culturales",
            "source": SOURCE_URL,
            "records": len(events) + len(variable) + len(undated),
            "dated_records": len(events),
            "variable_date_records": len(variable),
            "undated_records": len(undated),
            "synced_from_onedrive": True,
        },
        "events": events,
        "variable_events": variable,
        "undated_records": undated,
        "issues": issues,
    }


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        xlsx = Path(td) / "Efemerides_cientificas.xlsx"
        download(xlsx)
        payload = build(xlsx)

    current = None
    if OUTPUT.exists():
        try:
            current = json.loads(OUTPUT.read_text(encoding="utf-8"))
        except Exception:
            current = None

    same = (
        current
        and current.get("events") == payload.get("events")
        and current.get("variable_events") == payload.get("variable_events")
        and current.get("undated_records") == payload.get("undated_records")
    )
    if same:
        print("Sin cambios en efemérides científicas.")
        return

    payload["meta"]["last_sync_utc"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    OUTPUT.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Actualizado {OUTPUT}: {payload['meta']['records']} registros.")


if __name__ == "__main__":
    main()
